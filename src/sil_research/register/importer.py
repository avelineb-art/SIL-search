"""Manual import of a downloaded NDIS Provider Register export (CSV or
Excel) into a versioned snapshot.

Deliberately not automated against the register's own search interface -
see the build spec's legal/compliance constraints. The operator downloads
the file themselves and supplies its path and the date it was downloaded;
everything else here is parsing and validation.
"""

from __future__ import annotations

import csv
import re
from datetime import date, datetime
from pathlib import Path
from typing import Any

from sqlalchemy.orm import Session

from sil_research.config import get_register_columns
from sil_research.database import ProviderRegisterSnapshot, RegisterEntry
from sil_research.extraction.organisation import normalise_name
from sil_research.register.base import ColumnMapping, ImportResult, RegisterImportError

REQUIRED_LOGICAL_FIELDS = ["abn", "entity_name"]
_GROUP_SPLIT_PATTERN = re.compile(r"[;|,]")
_DATE_FORMATS = ["%Y-%m-%d", "%d/%m/%Y", "%d-%m-%Y", "%d %B %Y", "%d %b %Y"]


def _normalise_header(header: str) -> str:
    return re.sub(r"[^a-z0-9]+", "", header.strip().lower())


def resolve_columns(headers: list[str], column_aliases: dict[str, list[str]]) -> ColumnMapping:
    """Map each logical field to whichever actual header matches one of its
    configured aliases. Raises loudly, listing the file's real headers and
    what was expected, if a required field can't be resolved.
    """
    normalised_headers = {_normalise_header(h): h for h in headers if h}
    logical_to_actual: dict[str, str] = {}
    for logical, aliases in column_aliases.items():
        for alias in aliases:
            key = _normalise_header(alias)
            if key in normalised_headers:
                logical_to_actual[logical] = normalised_headers[key]
                break

    missing_required = [f for f in REQUIRED_LOGICAL_FIELDS if f not in logical_to_actual]
    if missing_required:
        expected = {f: column_aliases.get(f, []) for f in missing_required}
        raise RegisterImportError(
            "Register file is missing required column(s): "
            f"{missing_required}. Actual headers found in the file: {headers}. "
            f"Expected one of these aliases per missing field: {expected}. "
            "The register export's structure may have changed since config/register_columns.yml "
            "was written - add the real header name to the relevant list there rather than "
            "guessing at a fixed column position."
        )
    return ColumnMapping(logical_to_actual=logical_to_actual, all_headers=list(headers))


def _read_csv(path: Path) -> tuple[list[str], list[dict[str, Any]]]:
    with path.open(newline="", encoding="utf-8-sig") as fh:
        reader = csv.DictReader(fh)
        headers = list(reader.fieldnames or [])
        rows = [dict(row) for row in reader]
    return headers, rows


def _read_xlsx(path: Path) -> tuple[list[str], list[dict[str, Any]]]:
    from openpyxl import load_workbook

    workbook = load_workbook(path, read_only=True, data_only=True)
    worksheet = workbook.active
    rows_iter = worksheet.iter_rows(values_only=True)
    try:
        header_row = next(rows_iter)
    except StopIteration:
        return [], []
    headers = [str(cell).strip() if cell is not None else "" for cell in header_row]
    rows = []
    for raw_row in rows_iter:
        rows.append({headers[i]: raw_row[i] if i < len(raw_row) else None for i in range(len(headers))})
    return headers, rows


def _read_rows(path: Path) -> tuple[list[str], list[dict[str, Any]]]:
    if not path.exists():
        raise RegisterImportError(f"Register file not found: {path}")
    if path.suffix.lower() in (".xlsx", ".xlsm"):
        return _read_xlsx(path)
    return _read_csv(path)


def _clean(value: Any) -> str | None:
    if value is None:
        return None
    text = str(value).strip()
    return text or None


def _split_groups(value: Any) -> list[str]:
    text = _clean(value)
    if not text:
        return []
    return [part.strip() for part in _GROUP_SPLIT_PATTERN.split(text) if part.strip()]


def _parse_date(value: Any) -> datetime | None:
    if value is None or value == "":
        return None
    if isinstance(value, datetime):
        return value
    if isinstance(value, date):
        return datetime.combine(value, datetime.min.time())
    text = str(value).strip()
    for fmt in _DATE_FORMATS:
        try:
            return datetime.strptime(text, fmt)
        except ValueError:
            continue
    return None


def build_entry_data(raw_row: dict[str, Any], mapping: ColumnMapping) -> dict[str, Any] | None:
    """Pure transform from one raw row to RegisterEntry field values, or
    None if the row has neither an ABN nor an entity name (e.g. a blank
    trailing row) and so carries nothing to match against.
    """
    abn_raw = raw_row.get(mapping.logical_to_actual.get("abn", ""))
    abn_digits = re.sub(r"\D", "", str(abn_raw)) if abn_raw is not None else ""
    entity_name = _clean(raw_row.get(mapping.logical_to_actual.get("entity_name", "")))

    if not abn_digits and not entity_name:
        return None

    trading_name = None
    if "trading_name" in mapping.logical_to_actual:
        trading_name = _clean(raw_row.get(mapping.logical_to_actual["trading_name"]))

    registration_status = None
    if "registration_status" in mapping.logical_to_actual:
        registration_status = _clean(raw_row.get(mapping.logical_to_actual["registration_status"]))

    registration_groups: list[str] = []
    if "registration_groups" in mapping.logical_to_actual:
        registration_groups = _split_groups(raw_row.get(mapping.logical_to_actual["registration_groups"]))

    registration_expiry = None
    if "registration_expiry" in mapping.logical_to_actual:
        registration_expiry = _parse_date(raw_row.get(mapping.logical_to_actual["registration_expiry"]))

    state = None
    if "state" in mapping.logical_to_actual:
        state = _clean(raw_row.get(mapping.logical_to_actual["state"]))

    return {
        "abn": abn_digits or None,
        "entity_name": entity_name,
        "normalised_entity_name": normalise_name(entity_name) if entity_name else None,
        "trading_name": trading_name,
        "normalised_trading_name": normalise_name(trading_name) if trading_name else None,
        "registration_status": registration_status,
        "registration_groups": registration_groups,
        "registration_expiry": registration_expiry,
        "state": state,
        "raw_row": {k: (v if isinstance(v, (str, int, float, bool)) or v is None else str(v)) for k, v in raw_row.items()},
    }


def import_register_file(
    file_path: str | Path,
    snapshot_date: date,
    session: Session,
    column_aliases: dict[str, list[str]] | None = None,
) -> ImportResult:
    """Parse and store one register export as a new versioned snapshot.

    Never mutates or replaces a prior snapshot - each import is additive,
    so matches can always be re-run against an older snapshot if needed.
    """
    path = Path(file_path)
    headers, raw_rows = _read_rows(path)
    if not headers:
        raise RegisterImportError(f"Could not read a header row from {path} - is the file empty or malformed?")

    aliases = column_aliases or get_register_columns()
    mapping = resolve_columns(headers, aliases)

    snapshot = ProviderRegisterSnapshot(
        source_file=str(path),
        download_date=datetime.combine(snapshot_date, datetime.min.time()),
        row_count=len(raw_rows),
        column_manifest={"logical_to_actual": mapping.logical_to_actual, "all_headers": mapping.all_headers},
    )
    session.add(snapshot)
    session.flush()

    imported = 0
    skipped = 0
    warnings: list[str] = []
    for raw_row in raw_rows:
        fields = build_entry_data(raw_row, mapping)
        if fields is None:
            skipped += 1
            continue
        session.add(RegisterEntry(snapshot_id=snapshot.snapshot_id, **fields))
        imported += 1

    if skipped:
        warnings.append(f"{skipped} row(s) had neither an ABN nor an entity name and were skipped.")

    return ImportResult(
        snapshot_id=snapshot.snapshot_id,
        source_file=str(path),
        download_date=snapshot_date,
        row_count=len(raw_rows),
        imported=imported,
        skipped=skipped,
        column_mapping=mapping.logical_to_actual,
        warnings=warnings,
    )
